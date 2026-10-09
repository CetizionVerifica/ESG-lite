"""Build the PCF-0 pilot golden spreadsheet and its JSON twin.

Run:  python3 docs/pcf/phase0/build_golden.py
Writes pilot-golden.xlsx (live formulas) and pilot-golden.json (inputs +
expected outputs, computed here in plain Python the same way the sheet does).
The E1 golden test reads the JSON and must match `expected` within ±0.5%.

Every number below is a PLACEHOLDER chosen to look plausible. None of it is
Midal data. Replace the INPUT blocks with Midal's BOM and approved plant data
in phase 0, re-run, and commit both files together.
"""
import json
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
PLACEHOLDER = "PLACEHOLDER (not Midal data)"

# ---------------------------------------------------------------- inputs ---
STUDY = {
    "product": "EC-grade aluminium wire rod 9.5 mm",
    "company": "Midal Cables (pilot)",
    "site": "Bahrain plant (placeholder)",
    "declared_unit": "kg",
    "declared_unit_qty": 1,
    "mass_per_unit_kg": 1.0,
    "reference_start": "2025-01-01",
    "reference_end": "2025-12-31",
    "year_type": "CY",
    "boundary": "cradle_to_gate",
    "standard": "iso14067",
    "pcr_tag": "PEP PCR-ed4 / EN 50693 (optional)",
    "allocation_key": "mass",
    "cut_off_rule_pct": 1.0,
    "cut_off_max_total_pct": 5.0,
    "gwp_set": "AR6",
}

# id, name, material_group, unit, value kgCO2e/unit, source, licence
FACTORS = [
    ("F-AL-P", "Primary aluminium, molten metal, GCC", "aluminium", "kg", 8.60, PLACEHOLDER, "open"),
    ("F-AL-R", "Recycled aluminium (cut-off: remelt only)", "aluminium", "kg", 0.50, PLACEHOLDER, "open"),
    ("F-ALB", "Al-B master alloy (AlB4)", "aluminium", "kg", 12.00, PLACEHOLDER, "open"),
    ("F-LUB", "Rolling emulsion concentrate", "chemical", "kg", 1.80, PLACEHOLDER, "open"),
    ("F-PAL", "Wooden pallet (fossil part only)", "packaging", "kg", 0.25, PLACEHOLDER, "open"),
    ("F-STR", "Steel strapping", "packaging", "kg", 2.30, PLACEHOLDER, "open"),
    ("F-PE", "PE stretch film", "packaging", "kg", 2.60, PLACEHOLDER, "open"),
    ("F-LDF", "Industrial waste to landfill", "waste", "kg", 0.50, PLACEHOLDER, "open"),
    ("F-HAZ", "Spent emulsion to treatment", "waste", "kg", 0.90, PLACEHOLDER, "open"),
    # Real rows already in ESG-lite (docs/factor-files/UPLOAD-transport-factors.xlsx, 2025)
    ("F-ROAD", "Road - HGV (all diesel) - All rigids [tonne.km]", "transport", "tonne.km", 0.2408,
     "DEFRA 2024 Freighting goods + WTT (existing EmissionFactor row, 2025)", "open"),
    ("F-SEA", "Sea - Cargo ship - General Cargo", "transport", "tonne.km", 0.0162,
     "DEFRA 2024 Freighting goods + WTT (existing EmissionFactor row, 2025)", "open"),
]

# Materials, packaging, waste per declared unit (1 kg rod).
# id, stage, name, factor, recycled factor (or None), qty, unit, recycled %, data_type, dqr (tech, geo, time)
INPUTS = [
    ("I1", "A1", "Aluminium (incl. 1.2% process loss)", "F-AL-P", "F-AL-R", 1.012, "kg", 5.0, "secondary", (2, 2, 1)),
    ("I2", "A1", "Al-B master alloy", "F-ALB", None, 0.0020, "kg", 0.0, "secondary", (2, 3, 2)),
    ("I3", "A1", "Rolling emulsion concentrate", "F-LUB", None, 0.0008, "kg", 0.0, "secondary", (3, 3, 2)),
    ("I4", "A3_packaging", "Wooden pallet", "F-PAL", None, 0.0040, "kg", 0.0, "secondary", (2, 2, 2)),
    ("I5", "A3_packaging", "Steel strapping", "F-STR", None, 0.0015, "kg", 0.0, "secondary", (2, 2, 2)),
    ("I6", "A3_packaging", "PE stretch film", "F-PE", None, 0.0006, "kg", 0.0, "secondary", (2, 2, 2)),
    ("I7", "A3_waste", "Industrial waste to landfill", "F-LDF", None, 0.0030, "kg", 0.0, "primary", (2, 2, 1)),
    ("I8", "A3_waste", "Spent emulsion to treatment", "F-HAZ", None, 0.0005, "kg", 0.0, "primary", (2, 2, 1)),
]

# Inbound legs (A2). mass is the carried input's qty in tonnes (from INPUTS).
# id, name, carried input, mode factor, distance km, data_type, dqr
LEGS = [
    ("T1", "Molten metal, smelter to plant (road)", "I1", "F-ROAD", 5, "primary", (1, 1, 1)),
    ("T2", "Master alloy, Rotterdam to Khalifa Bin Salman port (sea)", "I2", "F-SEA", 11800, "secondary", (2, 2, 2)),
    ("T3", "Master alloy, port to plant (road)", "I2", "F-ROAD", 40, "secondary", (2, 2, 2)),
    ("T4", "Emulsion concentrate, Jebel Ali to plant (sea)", "I3", "F-SEA", 900, "secondary", (2, 2, 2)),
]

# Approved Scope 1 + 2 Emission totals at the site for the reference period (tCO2e).
# category, scope, period total tCO2e
PLANT_EMISSIONS = [
    ("Stationary combustion - Natural gas", 1, 18000.0),
    ("Mobile combustion - Diesel (forklifts)", 1, 600.0),
    ("Stationary combustion - LPG", 1, 150.0),
    ("Purchased electricity (location-based)", 2, 42000.0),
]
# Approved ProductionData at the site for the period (tonnes); the pilot product is first.
PRODUCTION = [
    ("EC-grade aluminium wire rod 9.5 mm", 120000.0),
    ("Alloy rod 6101", 60000.0),
    ("AAAC / ACSR conductor", 90000.0),
    ("Other products", 30000.0),
]
ALLOC_DQR = (1, 1, 1)  # measured plant data

STAGES = ["A1", "A2", "A3_energy", "A3_packaging", "A3_waste"]


# ----------------------------------------------------------- calculation ---
def compute():
    """Plain-Python twin of computePcf for the pilot; also the JSON expected values."""
    f = {r[0]: r[4] for r in FACTORS}
    qty = {r[0]: r[5] for r in INPUTS}
    items = []  # every line that carries emissions: inputs, legs, allocations
    for iid, stage, name, fid, rfid, q, unit, r, dt, dqr in INPUTS:
        rf = f[rfid] if rfid else f[fid]
        kg = q * ((1 - r / 100) * f[fid] + (r / 100) * rf)
        items.append({"id": iid, "stage": stage, "name": name, "kg": kg, "data_type": dt, "dqr": dqr, "cut_off_candidate": True})
    for tid, name, carried, fid, km, dt, dqr in LEGS:
        mass_t = qty[carried] / 1000  # carried kg per declared unit -> t
        kg = mass_t * km * f[fid]
        items.append({"id": tid, "stage": "A2", "name": name, "kg": kg, "data_type": dt, "dqr": dqr, "cut_off_candidate": True})
    product_t = PRODUCTION[0][1]
    site_t = sum(p[1] for p in PRODUCTION)
    share = product_t / site_t
    product_units = product_t * 1000 / STUDY["mass_per_unit_kg"]
    for n, (cat, scope, tco2e) in enumerate(PLANT_EMISSIONS, 1):
        kg = tco2e * 1000 * share / product_units
        items.append({"id": f"A3E{n}", "stage": "A3_energy", "name": cat, "kg": kg, "data_type": "primary", "dqr": ALLOC_DQR, "cut_off_candidate": False})

    total = sum(i["kg"] for i in items)
    by_stage = {s: sum(i["kg"] for i in items if i["stage"] == s) for s in STAGES}
    below = [i for i in items if i["cut_off_candidate"] and i["kg"] / total * 100 < STUDY["cut_off_rule_pct"]]
    below_pct = sum(i["kg"] for i in below) / total * 100
    primary = sum(i["kg"] for i in items if i["data_type"] == "primary") / total * 100
    dqr_dims = [sum(i["kg"] * i["dqr"][d] for i in items) / total for d in range(3)]
    dqr_overall = sum(dqr_dims) / 3
    plant_total = sum(p[2] for p in PLANT_EMISSIONS)
    covered = by_stage["A3_energy"] * product_units / 1000  # tCO2e of S1+S2 covered by this product's PCF
    return {
        "allocation": {"share_pct": share * 100, "product_output_units": product_units, "site_total_t": site_t},
        "total_kg_per_unit": total,
        "by_stage": by_stage,
        "by_input": {i["id"]: i["kg"] for i in items},
        "lines": [{"id": i["id"], "stage": i["stage"], "kgco2e_per_unit": i["kg"], "data_type": i["data_type"],
                   "cut_off_candidate": i["cut_off_candidate"]} for i in items],
        "cut_off": {"below_threshold_ids": [i["id"] for i in below], "below_threshold_total_pct": below_pct,
                    "within_limit": below_pct < STUDY["cut_off_max_total_pct"]},
        "primary_data_share_pct": primary,
        "dqr": {"technology": dqr_dims[0], "geography": dqr_dims[1], "time": dqr_dims[2], "overall": dqr_overall},
        # With one footprinted product, coverage equals its allocation share by construction;
        # it becomes an independent check once several products have footprints.
        "reconciliation": {"plant_s1_s2_tco2e": plant_total, "covered_tco2e": covered,
                           "coverage_pct": covered / plant_total * 100,
                           "note": "single product: coverage equals its allocation share by construction"},
    }


# ------------------------------------------------------------- workbook ----
BOLD = Font(bold=True)
HEAD = PatternFill("solid", fgColor="DDE7F0")
INPUT = PatternFill("solid", fgColor="FFF4CC")
WARN = Font(bold=True, color="B00020")


def header(ws, row, cols):
    for c, v in enumerate(cols, 1):
        cell = ws.cell(row=row, column=c, value=v)
        cell.font = BOLD
        cell.fill = HEAD


def widths(ws, ws_widths):
    for i, w in enumerate(ws_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def build_workbook(path):
    wb = Workbook()

    # README
    ws = wb.active
    ws.title = "README"
    lines = [
        ("PCF-0 golden spreadsheet · Midal pilot", BOLD),
        ("ALL NUMBERS ARE PLACEHOLDERS. This is NOT Midal data and must never be shown to a client as such.", WARN),
        ("Transport factors F-ROAD and F-SEA are the real DEFRA 2024 rows already in ESG-lite; everything else is invented.", None),
        ("", None),
        ("Product: " + STUDY["product"] + " · declared unit 1 kg · cradle-to-gate · ISO 14067 · AR6 GWP100", None),
        ("Reference period: " + STUDY["reference_start"] + " to " + STUDY["reference_end"] + " (CY) · allocation key: mass", None),
        ("", None),
        ("Yellow cells are inputs. Everything else is a formula; change an input and the Result sheet updates.", None),
        ("Sheets: Study, Factors, Inputs (A1, packaging, waste), Transport (A2), Allocation (A3 energy), Result, Reconciliation.", None),
        ("The calculation mirrors computePcf in E1 (docs/pcf/foundation/E1-data-and-engine/CLAUDE.md, steps 1-7).", None),
        ("pilot-golden.json carries the same inputs and the expected outputs for the E1 golden test (±0.5%).", None),
        ("Regenerate both with: python3 docs/pcf/phase0/build_golden.py", None),
    ]
    for r, (text, font) in enumerate(lines, 1):
        c = ws.cell(row=r, column=1, value=text)
        if font:
            c.font = font
    widths(ws, [120])

    # Study
    ws = wb.create_sheet("Study")
    header(ws, 1, ["Field", "Value"])
    for r, (k, v) in enumerate(STUDY.items(), 2):
        ws.cell(row=r, column=1, value=k)
        c = ws.cell(row=r, column=2, value=v)
        c.fill = INPUT
    study_row = {k: r for r, k in enumerate(STUDY.keys(), 2)}
    widths(ws, [24, 44])

    def study(k):
        return f"Study!$B${study_row[k]}"

    # Factors
    ws = wb.create_sheet("Factors")
    header(ws, 1, ["factor_id", "name", "material_group", "unit", "value_kgco2e", "gwp_set", "source", "licence"])
    frow = {}
    for r, (fid, name, grp, unit, val, src, lic) in enumerate(FACTORS, 2):
        for c, v in enumerate([fid, name, grp, unit, val, "AR6", src, lic], 1):
            cell = ws.cell(row=r, column=c, value=v)
            if c == 5:
                cell.fill = INPUT
        frow[fid] = r
    widths(ws, [10, 48, 14, 10, 14, 8, 58, 8])

    def fval(fid):
        return f"Factors!$E${frow[fid]}"

    # Inputs (A1, A3 packaging, A3 waste)
    ws = wb.create_sheet("Inputs")
    header(ws, 1, ["id", "stage", "name", "factor_id", "recycled_factor_id", "qty_per_unit", "unit",
                   "recycled_share_pct", "virgin_factor", "recycled_factor", "kgco2e_per_unit",
                   "data_type", "dqr_technology", "dqr_geography", "dqr_time"])
    irow = {}
    for r, (iid, stage, name, fid, rfid, q, unit, rec, dt, dqr) in enumerate(INPUTS, 2):
        rf = rfid or fid
        vals = [iid, stage, name, fid, rfid or "", q, unit, rec, f"={fval(fid)}", f"={fval(rf)}",
                f"=F{r}*((1-H{r}/100)*I{r}+(H{r}/100)*J{r})", dt, *dqr]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=c, value=v)
            if c in (6, 8, 13, 14, 15):
                cell.fill = INPUT
        irow[iid] = r
    in_last = 1 + len(INPUTS)
    widths(ws, [5, 13, 38, 9, 12, 12, 6, 12, 12, 12, 15, 10, 9, 9, 9])

    # Transport (A2)
    ws = wb.create_sheet("Transport")
    header(ws, 1, ["id", "stage", "name", "carried_input", "mass_t_per_unit", "factor_id", "distance_km",
                   "factor_kgco2e_per_tkm", "kgco2e_per_unit", "data_type", "dqr_technology", "dqr_geography", "dqr_time"])
    for r, (tid, name, carried, fid, km, dt, dqr) in enumerate(LEGS, 2):
        vals = [tid, "A2", name, carried, f"=Inputs!F{irow[carried]}/1000", fid, km, f"={fval(fid)}",
                f"=E{r}*G{r}*H{r}", dt, *dqr]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=c, value=v)
            if c in (7, 11, 12, 13):
                cell.fill = INPUT
    tr_last = 1 + len(LEGS)
    widths(ws, [5, 6, 52, 12, 15, 9, 11, 14, 15, 10, 9, 9, 9])

    # Allocation (A3 energy)
    ws = wb.create_sheet("Allocation")
    ws.cell(row=1, column=1, value="Approved ProductionData at the site in the reference period (t) · PLACEHOLDER, not Midal data").font = WARN
    header(ws, 2, ["product", "quantity_t"])
    for r, (p, t) in enumerate(PRODUCTION, 3):
        ws.cell(row=r, column=1, value=p)
        ws.cell(row=r, column=2, value=t).fill = INPUT
    p_last = 2 + len(PRODUCTION)
    k = p_last + 1
    ws.cell(row=k, column=1, value="site total (key_value_site_total)").font = BOLD
    ws.cell(row=k, column=2, value=f"=SUM(B3:B{p_last})")
    ws.cell(row=k + 1, column=1, value="pilot product (key_value_product)").font = BOLD
    ws.cell(row=k + 1, column=2, value="=B3")
    ws.cell(row=k + 2, column=1, value="share = product / site").font = BOLD
    ws.cell(row=k + 2, column=2, value=f"=B{k + 1}/B{k}")
    ws.cell(row=k + 3, column=1, value="product output in declared units").font = BOLD
    ws.cell(row=k + 3, column=2, value=f"=B{k + 1}*1000/{study('mass_per_unit_kg')}")
    share, units = f"$B${k + 2}", f"$B${k + 3}"
    h = k + 5
    ws.cell(row=h - 1, column=1, value="Approved Emission rows, Scope 1 + 2 only, same site and period (tCO2e)").font = BOLD
    header(ws, h, ["id", "category", "scope", "period_total_tco2e", "share", "allocated_kgco2e_per_unit",
                   "data_type", "dqr_technology", "dqr_geography", "dqr_time"])
    for n, (cat, scope, t) in enumerate(PLANT_EMISSIONS, 1):
        r = h + n
        vals = [f"A3E{n}", cat, scope, t, f"={share}", f"=D{r}*1000*E{r}/{units}", "primary", *ALLOC_DQR]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=c, value=v)
            if c == 4:
                cell.fill = INPUT
    al_first, al_last = h + 1, h + len(PLANT_EMISSIONS)
    widths(ws, [44, 42, 7, 18, 10, 24, 10, 9, 9, 9])

    # Result
    ws = wb.create_sheet("Result")
    ws.cell(row=1, column=1, value="Result per declared unit (kg CO2e / 1 kg rod) · PLACEHOLDER NUMBERS").font = WARN
    # by-line table: one formula row per input, leg, allocation
    header(ws, 3, ["id", "stage", "name", "kgco2e_per_unit", "share_of_total_pct", "below_cut_off",
                   "data_type", "dqr_technology", "dqr_geography", "dqr_time"])
    src = [("Inputs", r, "A", "B", "C", "K", "L", "M", "N", "O", True) for r in range(2, in_last + 1)]
    src += [("Transport", r, "A", "B", "C", "I", "J", "K", "L", "M", True) for r in range(2, tr_last + 1)]
    src += [("Allocation", r, "A", None, "B", "F", "G", "H", "I", "J", False) for r in range(al_first, al_last + 1)]
    first = 4
    last = first + len(src) - 1
    tot = f"$D${last + 3}"
    for n, (sh, r, cid, cst, cnm, ckg, cdt, c1, c2, c3, cut) in enumerate(src):
        row = first + n
        ws.cell(row=row, column=1, value=f"={sh}!{cid}{r}")
        ws.cell(row=row, column=2, value=f"={sh}!{cst}{r}" if cst else "A3_energy")
        ws.cell(row=row, column=3, value=f"={sh}!{cnm}{r}")
        ws.cell(row=row, column=4, value=f"={sh}!{ckg}{r}")
        ws.cell(row=row, column=5, value=f"=D{row}/{tot}*100")
        ws.cell(row=row, column=6, value=f'=IF(E{row}<{study("cut_off_rule_pct")},"yes","no")' if cut else "n/a")
        ws.cell(row=row, column=7, value=f"={sh}!{cdt}{r}")
        ws.cell(row=row, column=8, value=f"={sh}!{c1}{r}")
        ws.cell(row=row, column=9, value=f"={sh}!{c2}{r}")
        ws.cell(row=row, column=10, value=f"={sh}!{c3}{r}")
    rng = lambda col: f"{col}{first}:{col}{last}"  # noqa: E731
    s = last + 2
    rows = [
        ("by stage", None),
        ("TOTAL kg CO2e per declared unit", f"=SUM({rng('D')})"),
    ]
    for st in STAGES:
        rows.append((st, f'=SUMIF({rng("B")},"{st}",{rng("D")})'))
    rows += [
        ("", None),
        ("primary_data_share_pct", f'=SUMIF({rng("G")},"primary",{rng("D")})/{tot}*100'),
        ("dqr_technology (emission-weighted)", f"=SUMPRODUCT({rng('D')},{rng('H')})/{tot}"),
        ("dqr_geography (emission-weighted)", f"=SUMPRODUCT({rng('D')},{rng('I')})/{tot}"),
        ("dqr_time (emission-weighted)", f"=SUMPRODUCT({rng('D')},{rng('J')})/{tot}"),
        ("dqr_overall (mean of the three)", None),
        ("", None),
        ("cut-off: lines below threshold, % of total", f'=SUMIF({rng("F")},"yes",{rng("D")})/{tot}*100'),
        ("cut-off: within limit (< 5% in total)", None),
    ]
    for n, (label, formula) in enumerate(rows):
        r = s + n
        ws.cell(row=r, column=3, value=label).font = BOLD
        if formula:
            ws.cell(row=r, column=4, value=formula)
    assert f"$D${s + 1}" == tot
    dq = s + rows.index(("dqr_overall (mean of the three)", None))
    ws.cell(row=dq, column=4, value=f"=AVERAGE(D{dq - 3}:D{dq - 1})")
    co = s + len(rows) - 1
    ws.cell(row=co, column=4, value=f'=IF(D{co - 1}<{study("cut_off_max_total_pct")},"yes","NO")')
    widths(ws, [5, 13, 52, 16, 16, 12, 10, 9, 9, 9])

    # Reconciliation
    ws = wb.create_sheet("Reconciliation")
    ws.cell(row=1, column=1, value="Σ PCF(A3 energy) × approved production vs plant Scope 1+2 (only the pilot is footprinted, so coverage equals its allocation share by construction)").font = BOLD
    a3 = f"Result!$D${s + 1 + 1 + STAGES.index('A3_energy')}"
    recon = [
        ("plant Scope 1+2, tCO2e", f"=SUM(Allocation!D{al_first}:D{al_last})"),
        ("pilot A3 energy kg/unit", f"={a3}"),
        ("pilot output, declared units", f"=Allocation!{units}"),
        ("covered tCO2e", "=B4*B5/1000"),
        ("coverage %", "=B6/B3*100"),
        ("flag (< 90% = a product is missing a footprint or the key is off)", '=IF(B7<90,"below 90%","ok")'),
    ]
    for n, (label, f) in enumerate(recon, 3):
        ws.cell(row=n, column=1, value=label)
        ws.cell(row=n, column=2, value=f)
    widths(ws, [66, 18])

    # Banner on every sheet so no screenshot can pass for Midal data.
    for sheet in wb.worksheets[1:]:
        sheet.sheet_properties.tabColor = "B00020"
        col = sheet.max_column + 2
        sheet.cell(row=1, column=col, value="PLACEHOLDER NUMBERS · not Midal data").font = WARN
        sheet.column_dimensions[get_column_letter(col)].width = 40
    for sheet in wb.worksheets:
        for row in sheet.iter_rows():
            for c in row:
                c.alignment = Alignment(vertical="top")
    wb.save(path)
    return {"total_cell": f"Result!{tot.replace('$', '')}"}


def main():
    expected = compute()
    meta = build_workbook(os.path.join(HERE, "pilot-golden.xlsx"))
    doc = {
        "_note": "PLACEHOLDER numbers, not Midal data. Generated by build_golden.py; edit that file, not this one.",
        "tolerance_pct": 0.5,
        "study": STUDY,
        "factors": [dict(zip(["id", "name", "material_group", "unit", "value_kgco2e", "source", "licence"], f)) | {"gwp_set": "AR6"} for f in FACTORS],
        "inputs": [
            {"id": i[0], "stage": i[1], "name": i[2], "material_factor_id": i[3], "recycled_factor_id": i[4],
             "quantity": i[5], "unit": i[6], "recycled_share_pct": i[7], "data_type": i[8],
             "dqr_technology": i[9][0], "dqr_geography": i[9][1], "dqr_time": i[9][2]}
            for i in INPUTS
        ],
        "transport_legs": [
            {"id": t[0], "stage": "A2", "name": t[1], "carried_input_id": t[2],
             "mass_t_per_unit": next(i[5] for i in INPUTS if i[0] == t[2]) / 1000, "factor_id": t[3],
             "distance_km": t[4], "data_type": t[5], "dqr_technology": t[6][0], "dqr_geography": t[6][1], "dqr_time": t[6][2]}
            for t in LEGS
        ],
        "plant_emissions_approved": [{"category": c, "scope": s, "period_total_tco2e": t} for c, s, t in PLANT_EMISSIONS],
        "production_approved_t": [{"product": p, "quantity_t": t, "is_pilot": n == 0} for n, (p, t) in enumerate(PRODUCTION)],
        "allocation_dqr": dict(zip(["technology", "geography", "time"], ALLOC_DQR)),
        "expected": expected,
        "xlsx_total_cell": meta["total_cell"],
    }
    with open(os.path.join(HERE, "pilot-golden.json"), "w") as fh:
        json.dump(doc, fh, indent=2)
        fh.write("\n")
    print(f"total = {expected['total_kg_per_unit']:.6f} kg CO2e / kg  ({meta['total_cell']})")
    for k, v in expected["by_stage"].items():
        print(f"  {k:13s} {v:.6f}")


if __name__ == "__main__":
    main()
